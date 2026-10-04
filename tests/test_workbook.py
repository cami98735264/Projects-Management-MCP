from pathlib import Path

import openpyxl

from helpers import act, crash_activities, draft_from_arrows, load_fixture, normalized
from pm_mcp.domain.errors import IssueCode
from pm_mcp.domain.models import ProbabilityQuery, ProjectDraft, ProjectMetadata
from pm_mcp.domain.rules import effective_outputs
from pm_mcp.output.formula_eval import WorkbookEvaluator
from pm_mcp.output.workbook import SheetKind as K
from pm_mcp.output.workbook import plan_sheets
from pm_mcp.output.workbook_validation import validate_workbook
from pm_mcp.service import generate_workbook


def plan(**kw):
    project = normalized(ProjectDraft(**kw))
    return plan_sheets(project, effective_outputs(project, project.calculation_method))


def test_sheet_plan_follows_requested_outputs_and_data_shape():
    cpm = [act("A", [], 1), act("B", ["A"], 2)]
    pert = [act("A", est=(1, 2, 3)), act("B", ["A"], est=(1, 2, 4))]
    both = [act("A", [], 2, est=(1, 2, 3)), act("B", ["A"], 3, est=(1, 3, 4))]
    assert plan(activities=cpm, requested_outputs=["CRITICAL_PATH"]) == [K.PROBLEM, K.DATA, K.CPM, K.RESULTS]
    assert plan(activities=pert, probability_queries=[ProbabilityQuery(kind="AT_MOST", upper_bound=4)]) == [
        K.PROBLEM, K.DATA, K.PERT, K.CPM, K.PROBABILITY, K.RESULTS]
    assert plan(activities=pert, requested_outputs=["GANTT"]) == [K.PROBLEM, K.DATA, K.PERT, K.CPM, K.GANTT, K.RESULTS]
    assert K.CPM_DETERMINISTIC in plan(activities=both, requested_outputs=["EXPECTED_DURATION"])
    full_cpm = plan(activities=cpm, requested_outputs=["FULL_REPORT"])
    assert K.NETWORK in full_cpm and K.GANTT in full_cpm and K.PERT not in full_cpm and K.PROBABILITY not in full_cpm
    crash = crash_activities(load_fixture("crashing_reference_problems.json")["problem_1"])
    assert plan(activities=crash, requested_outputs=["CRASHING"]) == [K.PROBLEM, K.DATA, K.CRASHING, K.RESULTS]


def exercise_draft(**kw):
    return draft_from_arrows(load_fixture("taller1_transcription.json"), **kw)


def test_formulas_are_live_and_validated(output_dir):
    draft = exercise_draft(requested_outputs=["FULL_REPORT"],
                           probability_queries=[ProbabilityQuery(id="p", kind="AT_MOST", upper_bound=13),
                                                ProbabilityQuery(id="q", kind="PERCENTILE_TO_DURATION", target_probability=0.98)])
    result = generate_workbook(draft, "live.xlsx")
    report = result.workbook_validation
    assert result.ok, report.issues[:5]
    assert result.sheets == ["Enunciado", "Datos", "PERT", "CPM", "Red AON", "Gantt", "Probabilidad", "Resultados"]
    assert report.formula_count == report.evaluated_formula_count == report.expected_values_checked > 500
    assert report.conditional_format_rules > 0

    wb = openpyxl.load_workbook(result.path)
    cpm = wb["CPM"]
    rows = {cpm.cell(r, 1).value: r for r in range(6, 30) if cpm.cell(r, 1).value}
    assert cpm.cell(rows["B"], 9).value == f"=ROUND(G{rows['B']}-E{rows['B']},10)"
    assert cpm.cell(rows["F"], 5).value.startswith("=MAX(")

    data = wb["Datos"]
    header = [c.value for c in data[4]]
    b_col = header.index("Pesimista (b)") + 1
    j_row = next(r for r in range(5, 30) if data.cell(r, 1).value == "J")
    cells = {(ws.title, c.row, c.column): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if c.value is not None}
    cells[("Datos", j_row, b_col)] = 21  # t_e(J) = (1 + 21 + 4·5)/6 = 7 → J finishes at 10 + 7 = 17
    assert WorkbookEvaluator(cells, wb.sheetnames).value(("CPM", 2, 4)) == 17


def test_validator_detects_broken_workbooks(output_dir, tmp_path):
    result = generate_workbook(ProjectDraft(activities=[act("A", [], 2), act("B", ["A"], 3)], requested_outputs=["SLACK"]), "base.xlsx")
    assert result.ok
    wb = openpyxl.load_workbook(result.path)
    ws = wb["CPM"]
    ws["F6"] = "=#REF!+1"
    ws["G6"] = "='Nope'!A1"
    ws["H6"] = "=Z99+1"
    ws["I6"] = "=FOO(1)"
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    report = validate_workbook(broken, expected_sheets=[*result.sheets, "Missing"])
    kinds = {i.kind for i in report.issues}
    assert {"formula_error_token", "unknown_sheet_reference", "dangling_reference", "unknown_function", "missing_sheet"} <= kinds
    assert not report.passed


def test_output_is_reproducible_and_paths_are_guarded(output_dir):
    draft = ProjectDraft(activities=[act("A", [], 2), act("B", ["A"], 3)], requested_outputs=["GANTT"])
    first, second = generate_workbook(draft, "a.xlsx"), generate_workbook(draft, "b.xlsx")
    assert Path(first.path).read_bytes() == Path(second.path).read_bytes()
    again = generate_workbook(draft, "a.xlsx")
    assert not again.ok and again.issues[-1].code == IssueCode.OUTPUT_PATH
    assert generate_workbook(draft, "a.xlsx", overwrite=True).ok
    assert not generate_workbook(draft, "a.csv").ok


def test_english_workbook(output_dir):
    draft = ProjectDraft(metadata=ProjectMetadata(language="en"), activities=[act("A", [], 2), act("B", ["A"], 3)],
                         requested_outputs=["GANTT", "CRITICAL_PATH"])
    result = generate_workbook(draft, "en.xlsx")
    assert result.ok and result.sheets == ["Problem", "Activity Data", "CPM", "Gantt", "Final Results"]
    procedures = {a.outputs[0]: a.procedure for a in result.answers}
    assert "1. Critical activities (slack = 0)\n   A, B" in procedures["CRITICAL_PATH"]
    assert "A → B: Σ t = 2 + 3 = 5" in procedures["CRITICAL_PATH"]
    assert "⌈T⌉ = ⌈5⌉ = 5" in procedures["GANTT"]
