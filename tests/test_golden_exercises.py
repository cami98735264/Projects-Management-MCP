"""Golden end-to-end tests: the exercise document and the reference PERT example.

Expected values are independently derived (hand arithmetic in test_pert_and_cpm.py, scipy for the normal
distribution in test_probability.py); here they check the full document → answers → workbook chain.
"""

import math
from fractions import Fraction as F

import openpyxl

from helpers import arcs_from, draft_from_arrows, load_fixture
from pm_mcp.domain.models import AnalysisOptions, ProjectDraft, ProjectMetadata
from pm_mcp.engine.aoa import convert_arrow_network
from pm_mcp.ingestion.document_extraction import extract_document_content
from pm_mcp.ingestion.normalize import normalize_project_input
from pm_mcp.ingestion.requirements_inference import suggest_requirements_from_text
from pm_mcp.pipeline import solve_project
from pm_mcp.service import generate_workbook


def test_exercise_document_end_to_end(exercise_docx, output_dir, tmp_path):
    fixture = load_fixture("taller1_transcription.json")
    doc = extract_document_content(exercise_docx, tmp_path / "images")
    extracted = {img.part: img.sha256 for img in doc.images}
    for part, sha in fixture["source_images"].items():
        assert extracted[part] == sha, "the transcription must correspond to these exact images"

    suggestions = suggest_requirements_from_text("\n".join(doc.text_blocks))
    conversion = convert_arrow_network(arcs_from(fixture))
    draft = ProjectDraft(
        metadata=ProjectMetadata(title=fixture["title"], time_unit=suggestions.time_unit, language=suggestions.language,
                                 source_description=doc.path),
        activities=conversion.activities, arrow_network_trace=conversion.trace,
        questions=suggestions.questions, probability_queries=suggestions.probability_queries)
    normalization = normalize_project_input(draft)
    assert normalization.ok
    solution = solve_project(normalization.project)

    assert solution.validation.passed
    assert solution.project.calculation_method == "PERT"
    assert solution.cpm.project_duration == 15
    assert solution.cpm.critical_paths == [["A", "C", "F", "H", "J"]]
    assert solution.variance.total_variance == F(37, 9)
    probabilities = {r.query.id: r for r in solution.probabilities}
    assert math.isclose(probabilities["d"].probability_exact, 0.16196992279926703, rel_tol=1e-9)
    assert math.isclose(probabilities["e"].probability_exact, 0.31093671216537044, rel_tol=1e-9)
    assert math.isclose(probabilities["f"].probability_exact, 0.3041048557312799, rel_tol=1e-9)
    assert solution.percentiles[0].duration_whole_units == 20

    answers = {a.label: a.answer for a in solution.answers}
    assert list(answers) == list("abcdefg")
    assert "Ruta crítica: A → C → F → H → J" in answers["c"] and "σ² = 37/9" in answers["c"]
    assert answers["c"].startswith("• Tiempo esperado del proyecto: 15 semanas.")  # direct result first
    assert all(line.startswith(("•", "   ", "Tiempos")) for a in answers.values() for line in a.split("\n"))
    assert "0.1620" in answers["d"] and "0.3109" in answers["e"] and "0.3041" in answers["f"]
    assert "20 semanas" in answers["g"]

    procedures = {a.label: a.procedure for a in solution.answers}
    assert "1. Paso hacia adelante" in procedures["b"] and "\n   B (t = 3): IC = TC A = 2; TC = 2 + 3 = 5\n" in procedures["b"]
    assert "   F (t = 1): IC = máx(TC C = 6, TC B = 5) = 6; TC = 6 + 1 = 7" in procedures["b"]
    assert "   B (t = 3): TL = mín(IL D = 8, IL F = 6, IL G = 10) = 6; IL = 6 − 3 = 3" in procedures["b"]
    assert "3. Holgura = IL − IC" in procedures["b"] and "   B: 3 − 2 = 1\n" in procedures["b"]
    assert "C: (2 + 10 + 4·3) / 6 = 4" in procedures["c"] and "C: σ² = [(10 − 2) / 6]² = 16/9" in procedures["c"]
    assert "constantes: 4 = peso de m" in procedures["c"] and "2 + 4 + 1 + 3 + 5 = 15" in procedures["c"]
    assert "Z = (13 − 15) / 2.0276 = −0.9864" in procedures["d"]
    assert "1 − Φ(Z)" in procedures["e"] and "Φ(Z₂) − Φ(Z₁)" in procedures["f"]
    assert "T_p = t_e + Z · σ" in procedures["g"] and "⌈19.1642⌉ = 20 semanas" in procedures["g"]

    result = generate_workbook(draft, "solucion.xlsx", layout="classic")
    assert result.ok and result.workbook_validation.passed
    assert result.sheets == ["Enunciado", "Datos", "PERT", "CPM", "Gantt", "Probabilidad", "Resultados"]
    results = openpyxl.load_workbook(result.path)["Resultados"]
    label_rows = {results.cell(r, 1).value: r for r in range(5, 60) if results.cell(r, 1).value in set("abcdefg")}
    assert list(label_rows) == list("abcdefg")
    assert results.cell(4, 4).value == "Procedimiento (variables y operaciones)"
    d = label_rows["d"]
    assert results.cell(d, 3).value == answers["d"]
    steps = solution.answers[3].procedure_steps
    assert len(steps) == label_rows["e"] - d  # one row per procedure step
    assert str(results.cell(d, 4).value).startswith("1. P(T ≤ 13): Z = (T_p − t_e) / σ")
    assert results.cell(d, 5).value.startswith("='Probabilidad'!")


def test_reference_pert_example_end_to_end(output_dir):
    fixture = load_fixture("pert_example_transcription.json")
    s = suggest_requirements_from_text(fixture["question_text"])
    draft = draft_from_arrows(fixture, metadata=ProjectMetadata(title=fixture["title"], time_unit=s.time_unit),
                              questions=s.questions, probability_queries=s.probability_queries,
                              options=AnalysisOptions(probability_method="table"))
    solution = solve_project(normalize_project_input(draft).project)
    estimates = {e.activity_id: (e.expected_duration, e.variance) for e in solution.pert.estimates}
    # (a + b + 4m)/6 and ((b − a)/6)² per activity
    assert estimates == {"A": (3, F(4, 9)), "B": (3, F(4, 9)), "C": (5, F(1, 9)), "D": (4, 1), "E": (8, F(1, 9)),
                         "F": (6, F(4, 9)), "G": (5, F(1, 9)), "H": (9, F(4, 9)), "I": (3, F(4, 9))}
    assert solution.cpm.critical_paths == [["A", "E", "F", "G", "I"]] and solution.cpm.project_duration == 25
    assert solution.variance.total_variance == F(14, 9)
    assert [r.probability for r in solution.probabilities] == [0.9452, 0.5, 0.0548, 0.0548, 0.4452]
    percentile = solution.percentiles[0]
    assert percentile.z_table == 1.29 and round(percentile.duration, 4) == 26.6089 and percentile.duration_whole_units == 27

    answers = {a.label: a.answer for a in solution.answers}
    assert "A: t_e = 3, σ² = 4/9" in answers["b"] and "D: t_e = 4, σ² = 1" in answers["b"]

    result = generate_workbook(draft, "ejemplo.xlsx", layout="classic")
    assert result.ok and "Red AON" in result.sheets
