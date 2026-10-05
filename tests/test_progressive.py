"""Progressive (cumulative, one sheet per step) workbook layout.

Every expected number comes from the engine (the classic workbook's expectations); the hand-computed Taller 2
figures are asserted against the LibreOffice recalculation, independently of this code.
"""

from __future__ import annotations

import importlib.util
import json
import math
import re
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest

from helpers import crash_activities, draft_from_arrows, load_fixture, normalized
from pm_mcp.domain.models import (
    AnalysisOptions,
    CrashingRequest,
    IndirectCostPoint,
    ProbabilityQuery,
    ProjectDraft,
    ProjectMetadata,
)
from pm_mcp.ingestion.requirements_inference import suggest_requirements_from_text
from pm_mcp.output.formula_eval import WorkbookEvaluator
from pm_mcp.output.progressive import MAX_TAB, write_progressive_workbook
from pm_mcp.output.workbook_validation import validate_workbook, values_match
from pm_mcp.pipeline import solve_project
from pm_mcp.service import check_workbook, generate_workbook

DOUBLE = set("═║╔╗╚╝╠╣╦╩╬╒╕╘╛╞╡╤╧╪╓╖╙╜╟╢╥╨╫")
MERGE_PY = Path("/home/ubuntu/development/services/pm-solve-api/src/pm_solve_api/merge.py")

# literal narrative and questions of the Taller 2 statement (tests/fixtures/taller2_projects.json holds the data)
P1_CONTEXT = ("La constructora Andina planea concursar en un proyecto de construccion de un puente en el departamento "
              "del Huila. Al preparar sus estimaciones reunio los datos de la tabla. Todos los tiempos estan en semanas "
              "y los costos en dolares.")
P2_CONTEXT = ("Existe una multa de 600000 pesos si el proyecto no se entrega en 16 semanas. Por lo tanto, Maria Camila "
              "esta muy interesada en la probabilidad de tener el proyecto a tiempo, para lo que requiere que Usted:")


def with_network(project: dict) -> dict:
    return {**project, "requested_outputs": [*(project.get("requested_outputs") or []), "NETWORK_DIAGRAM"]}


def taller2(i: int, context: bool = True) -> ProjectDraft:
    project = with_network(load_fixture("taller2_projects.json")["projects"][i])
    if context:
        project["metadata"] = {**project["metadata"], "context": (P1_CONTEXT, P2_CONTEXT)[i]}
    return ProjectDraft.model_validate(project)


def taller1() -> ProjectDraft:
    return draft_from_arrows(load_fixture("taller1_transcription.json"), requested_outputs=["FULL_REPORT"],
                             probability_queries=[ProbabilityQuery(id="p", kind="AT_MOST", upper_bound=13),
                                                  ProbabilityQuery(id="q", kind="PERCENTILE_TO_DURATION",
                                                                   target_probability=0.98)])


def pert_example() -> ProjectDraft:
    fixture = load_fixture("pert_example_transcription.json")
    s = suggest_requirements_from_text(fixture["question_text"])
    return draft_from_arrows(fixture, metadata=ProjectMetadata(title=fixture["title"], time_unit=s.time_unit),
                             questions=s.questions, probability_queries=s.probability_queries,
                             requested_outputs=["NETWORK_DIAGRAM", "GANTT"],
                             options=AnalysisOptions(probability_method="table"))


def crashing(name: str) -> ProjectDraft:
    fixture = load_fixture("crashing_reference_problems.json")[name]
    request = CrashingRequest(indirect_costs=[IndirectCostPoint(**p) for p in fixture.get("indirect_costs", [])])
    return ProjectDraft(metadata=ProjectMetadata(title=name), activities=crash_activities(fixture),
                        requested_outputs=["CRASHING", "NETWORK_DIAGRAM", "GANTT", "CRITICAL_PATH"], crashing=request)


CASES = {
    "taller1": taller1,
    "pert_example": pert_example,
    "crash_problem_1": lambda: crashing("problem_1"),
    "crash_problem_2_indirect": lambda: crashing("problem_2"),
    "taller2_p1_penalty": lambda: taller2(0),
    "taller2_p2_pert": lambda: taller2(1),
    "english": lambda: taller2(1).model_copy(update={"metadata": ProjectMetadata(title="Bid", language="en")}),
}


def sheets(path, data_only=False):
    return openpyxl.load_workbook(path, data_only=data_only)


def cells_of(ws) -> dict[tuple[int, int], object]:
    return {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")}


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    """Every case generated once through the service (default layout) plus the raw progressive build."""
    out = tmp_path_factory.mktemp("progressive")
    import os

    os.environ["PM_MCP_OUTPUT_DIR"] = str(out)
    data = {}
    for name, make in CASES.items():
        draft = make()
        result = generate_workbook(draft, f"{name}.xlsx", overwrite=True)
        classic = generate_workbook(draft, f"{name}_classic.xlsx", overwrite=True, layout="classic")
        solution = solve_project(normalized(draft))
        build = write_progressive_workbook(solution, out / f"{name}_raw.xlsx")
        data[name] = (draft, result, classic, solution, build)
    return data


@pytest.mark.parametrize("name", list(CASES))
def test_progressive_is_the_default_and_validates(built, name):
    _, result, _, _, build = built[name]
    report = result.workbook_validation
    assert result.ok and result.layout == "progressive", report.issues[:5]
    assert not [n for n in result.notes if "classic layout delivered" in n], result.notes
    assert report.passed and report.formula_count == report.evaluated_formula_count == report.expected_values_checked
    assert result.sheets == build.sheet_names
    assert result.sheets[0].startswith("1.1 ") and result.sheets[1].startswith("1.2 ")
    for tab in result.sheets:
        assert len(tab) <= MAX_TAB and ":" not in tab and "/" not in tab
        assert len("P12 " + tab) <= 31, "merge.py's prefix must fit without truncation"
    numbers = [tab.split(" ")[0] for tab in result.sheets]
    assert numbers[2:] == [str(i) for i in range(2, len(numbers))] or numbers[2] == "1.3"


@pytest.mark.parametrize("name", list(CASES))
def test_every_sheet_is_cumulative_and_shows_nothing_ahead(built, name):
    _, result, _, _, build = built[name]
    wb = sheets(result.path)
    index = {s.tab: i for i, s in enumerate(build.steps)}
    previous: dict | None = None
    for ws in wb.worksheets:
        k = index[ws.title]
        here = cells_of(ws)
        for (r, c), value in here.items():
            if (r, c) in build.structural or r <= 2:
                continue
            assert (r, c) in build.first_step, (ws.title, r, c, value)
            assert build.first_step[(r, c)] <= k, f"{ws.title}: {value!r} belongs to a later step"
            assert build.tab_cells[ws.title][(r, c)] <= k
        if previous is not None:
            lost = [key for key in previous if key not in here and key not in build.structural and key[0] > 2]
            assert not lost, f"{ws.title} dropped cells of the previous sheet: {lost[:5]}"
        previous = here


@pytest.mark.parametrize("name", list(CASES))
def test_no_emoji_anywhere(built, name):
    import re

    emoji = re.compile("[\U0001F000-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F]")
    _, result, _, _, _ = built[name]
    for ws in sheets(result.path).worksheets:
        bad = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and emoji.search(c.value)]
        assert not bad, f"{ws.title}: {bad[:3]}"


@pytest.mark.parametrize("name", list(CASES))
def test_networks_are_on_screen_below_the_statement(built, name):
    """The AON network (and, with crashing, the network of every step) starts in the first columns and comes
    right after the statement, before the tables: off to the right it looked as if the nodes had been removed."""
    _, result, _, _, build = built[name]
    wb = sheets(result.path)
    last = wb.worksheets[-1]
    blocks = {b.key: b for b in build.blocks}
    if "net" not in blocks:
        pytest.skip("no network requested")
    net = blocks["net"]
    assert net.first_col == 1
    tables = [b for b in blocks.values() if b.lane == "left"]
    assert all(net.top < b.top for b in tables)
    tops = [b.top for b in blocks.values() if b.lane == "top"]
    assert all(t < net.top for t in tops)
    # node cards live within the first 40 columns of the final sheet
    ids = {a.id for a in built[name][3].project.activities if not a.is_dummy}
    card_cols = [int(c.column or 0) for row in last.iter_rows(min_row=net.top, max_row=net.top + net.max_row)
                 for c in row if c.value in ids]
    assert card_cols and max(card_cols) <= 40
    if "step_net" in blocks:
        step = blocks["step_net"]
        assert step.first_col == 1
        labels = [c.value for row in last.iter_rows(min_col=1, max_col=1) for c in row
                  if isinstance(c.value, str) and re.fullmatch(r"(PASO|STEP) \d+", c.value.strip().upper())]
        assert len(labels) == len(built[name][3].crashing.states)


@pytest.mark.parametrize("name", list(CASES))
def test_formulas_only_reference_their_own_sheet_and_each_sheet_stands_alone(built, name):
    _, result, _, _, build = built[name]
    wb, values = sheets(result.path), sheets(result.path, data_only=True)
    for ws in wb.worksheets:
        cells = {(ws.title, r, c): v for (r, c), v in cells_of(ws).items()}
        formulas = [key for key, v in cells.items() if isinstance(v, str) and v.startswith("=")]
        assert all("!" not in cells[key] for key in formulas), "a formula points to another sheet"
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                assert all("!" not in f for f in rule.formula or [])
        evaluator = WorkbookEvaluator(cells, [ws.title])  # this sheet only
        for key in formulas:
            got = evaluator.value(key)
            assert values_match(got, values[ws.title].cell(key[1], key[2]).value), (key, cells[key])
            if key in build.expected_values:
                assert values_match(got, build.expected_values[key])
        assert not evaluator.blank_references, evaluator.blank_references[:3]


@pytest.mark.parametrize("name", list(CASES))
def test_no_critical_marks_before_the_slack_step(built, name):
    _, result, _, _, build = built[name]
    keys = [s.key for s in build.steps]
    if "slack" not in keys:
        pytest.skip("no CPM step")
    slack = keys.index("slack")
    wb = sheets(result.path)
    for step in build.steps[:slack]:
        ws = wb[step.tab]
        text = "".join(str(v) for v in cells_of(ws).values())
        assert not DOUBLE & set(text), f"{step.tab} shows a critical (double) line"
        assert not list(ws.conditional_formatting), f"{step.tab} has conditional formats"
    later = "".join(str(v) for v in cells_of(wb[build.steps[slack].tab]).values())
    if any(n in build.classic.sheet_names for n in ("Red AON", "Network AON")):
        assert DOUBLE & set(later), "the slack step draws the critical path with double lines"
        assert list(wb[build.steps[slack].tab].conditional_formatting), "critical shading starts at the slack step"


@pytest.mark.parametrize("name", list(CASES))
def test_last_sheet_holds_every_answer_of_the_classic_results(built, name):
    _, result, classic, _, _ = built[name]
    results_name = classic.sheets[-1]

    def unsheet(v):  # sheet mentions differ on purpose (classic sheet → progressive tab)
        if not isinstance(v, str):
            return v
        v = re.sub(r"\bthe ('[^']*'|[\w ]+?) sheet\b", "sheet ·", v)
        return re.sub(r"\b(hoja|sheet) ('[^']*'|\w+)", r"\1 ·", v)

    expected = [unsheet(v) for v in cells_of(sheets(classic.path, data_only=True)[results_name]).values()]
    last = [unsheet(v) for v in cells_of(sheets(result.path, data_only=True)[result.sheets[-1]]).values()]
    missing = [v for v in expected if not any(values_match(v, w) for w in last)]
    assert not missing, missing[:5]


@pytest.mark.parametrize("name", ["crash_problem_1", "crash_problem_2_indirect", "taller2_p1_penalty"])
def test_one_sheet_per_engine_compression_step(built, name):
    _, result, _, solution, build = built[name]
    steps = len(solution.crashing.states) - 1
    crash_tabs = [s.tab for s in build.steps if s.key.startswith("crash:")]
    assert len(crash_tabs) == steps > 0
    durations = [st.project_duration for st in solution.crashing.states]
    for tab, (a, b) in zip(crash_tabs, zip(durations, durations[1:])):
        assert tab.endswith(f"{int(a)}-{int(b)}") if a == int(a) and b == int(b) else True
    keys = [s.key for s in build.steps]
    assert keys.index("slopes") < keys.index("crash:1")
    if solution.crashing.optimal is not None:
        assert keys.index(f"crash:{steps}") < keys.index("total") < keys.index("decision")
        # indirect / total panel rows of the step networks stay hidden until the total-cost step
        wb = sheets(result.path)
        total = keys.index("total")
        label = "Costos indirectos" if solution.project.metadata.language == "es" else "Indirect cost"
        before = [s.tab for s in build.steps[:total]]
        assert not any(isinstance(v, str) and v.startswith(label) for tab in before for v in cells_of(wb[tab]).values())


def test_context_and_literal_questions_on_the_first_sheets(built):
    draft, result, _, _, _ = built["taller2_p1_penalty"]
    wb = sheets(result.path)
    first = cells_of(wb[result.sheets[0]]).values()
    assert P1_CONTEXT in first and draft.metadata.title in first
    third = list(cells_of(wb[result.sheets[2]]).values())
    assert all(q.text in third for q in draft.questions)
    # without a context, only the title (and the time unit): nothing is made up
    no_ctx = taller2(0, context=False)
    build = write_progressive_workbook(solve_project(normalized(no_ctx)), Path(result.path).with_name("no_ctx.xlsx"))
    values = list(cells_of(sheets(build.path)[build.sheet_names[0]]).values())
    assert P1_CONTEXT not in values and no_ctx.metadata.title in values


def test_check_workbook_regenerates_the_same_layout(built):
    draft, result, classic, _, _ = built["taller2_p2_pert"]
    assert check_workbook(result.path, draft).passed
    assert check_workbook(classic.path, draft).passed


def test_progressive_steps_of_a_pert_problem(built):
    _, result, _, _, build = built["taller2_p2_pert"]
    keys = [s.key for s in build.steps]
    assert keys == ["ctx", "data", "questions", "pert", "net", "fwd", "bwd", "slack", "variance", "prob", "pct",
                    "results"]
    assert result.sheets[3] == "2 Tiempos esperados"
    wb = sheets(result.path)
    pert_tab = cells_of(wb[result.sheets[3]]).values()
    assert "¿Crítica?" not in pert_tab and not any(str(v).startswith("Varianza del proyecto") for v in pert_tab)


def _merge_module():
    if not MERGE_PY.is_file():
        pytest.skip("pm-solve-api merge.py not present")
    spec = importlib.util.spec_from_file_location("pm_merge", MERGE_PY)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_merge_two_progressive_problems(built, tmp_path):
    merge = _merge_module()
    parts = [Path(built[n][1].path) for n in ("taller2_p1_penalty", "taller2_p2_pert")]
    target = tmp_path / "merged.xlsx"
    names = merge.merge_workbooks(parts, target)
    expected_names = [f"P1 {n}" for n in built["taller2_p1_penalty"][1].sheets] + \
                     [f"P2 {n}" for n in built["taller2_p2_pert"][1].sheets]
    assert names == expected_names
    wb = sheets(target)
    for k, name in ((1, "taller2_p1_penalty"), (2, "taller2_p2_pert")):
        expected = built[name][1]
        source = sheets(expected.path, data_only=True)
        for tab in expected.sheets:
            ws = wb[f"P{k} {tab}"]
            cells = {(ws.title, rr, cc): v for (rr, cc), v in cells_of(ws).items()}
            evaluator = WorkbookEvaluator(cells, [ws.title])
            for key, v in cells.items():
                if isinstance(v, str) and v.startswith("="):
                    assert values_match(evaluator.value(key), source[tab].cell(key[1], key[2]).value)
            assert not evaluator.blank_references


def test_merged_file_keeps_cached_results(built, tmp_path):
    """openpyxl writes formulas without results; merge.py injects the result each source had cached, so a reader
    that does not recalculate (data_only, /ask) still sees every number."""
    merge = _merge_module()
    parts = [Path(built[n][1].path) for n in ("taller2_p1_penalty", "taller2_p2_pert")]
    target = tmp_path / "merged.xlsx"
    merge.merge_workbooks(parts, target)
    formulas, values = sheets(target), sheets(target, data_only=True)
    checked = 0
    for k, name in ((1, "taller2_p1_penalty"), (2, "taller2_p2_pert")):
        source = sheets(built[name][1].path, data_only=True)
        for tab in built[name][1].sheets:
            ws = formulas[f"P{k} {tab}"]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        expected = source[tab][c.coordinate].value
                        got = values[ws.title][c.coordinate].value
                        assert values_match(got, expected), (ws.title, c.coordinate, got, expected)
                        checked += 1
    assert checked > 3000


@pytest.mark.parametrize("name", list(CASES))
def test_texts_name_progressive_tabs_not_classic_sheets(built, name):
    """Fixed texts ("ver hoja 'Compresión'", "la hoja CPM") point at the tab where that content first appears."""
    path = built[name][1].path
    wb = sheets(path)
    tabs = set(wb.sheetnames)
    mentioned = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or c.value.startswith("="):
                    continue
                assert not re.search(r"\b(hoja|sheet) (?!')", c.value), (ws.title, c.value)
                for quoted in re.findall(r"(?:hoja|sheet) '([^']+)'", c.value):
                    assert quoted in tabs, (ws.title, c.value)
                    mentioned.add(quoted)
    assert mentioned


def _soffice() -> str | None:
    return shutil.which("soffice") or shutil.which("libreoffice")


@pytest.mark.skipif(_soffice() is None, reason="LibreOffice not installed")
def test_libreoffice_recalculates_taller2_with_the_hand_results(built, tmp_path):
    """Strip the cached values (openpyxl does not keep them), let LibreOffice recalculate every formula of every
    sheet, and compare with the engine and with the hand calculation of the Taller 2."""
    recalculated = {}
    for name in ("taller2_p1_penalty", "taller2_p2_pert"):
        result, build = built[name][1], built[name][4]
        stripped = tmp_path / f"{name}.xlsx"
        openpyxl.load_workbook(result.path).save(stripped)
        out = tmp_path / "lo"
        subprocess.run([_soffice(), "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(stripped)],
                       check=True, capture_output=True, timeout=240)
        formulas, values = sheets(stripped), sheets(out / stripped.name, data_only=True)
        checked = 0
        for ws in formulas.worksheets:
            for (r, c), v in cells_of(ws).items():
                if isinstance(v, str) and v.startswith("="):
                    got = values[ws.title].cell(r, c).value
                    assert got is not None and not (isinstance(got, str) and got.startswith("#")), (ws.title, r, c, v)
                    want = build.expected_values.get((ws.title, r, c))
                    if want is not None and want != "":
                        assert values_match(got, want) or (isinstance(got, float) and math.isclose(got, float(want),
                                                                                                    rel_tol=1e-6)), \
                            (ws.title, r, c, v, got, want)
                    checked += 1
        assert checked == len(build.expected_values)
        recalculated[name] = [v for v in cells_of(values[result.sheets[-1]]).values()]

    def has(name, number, places=4):
        return any(isinstance(v, (int, float)) and round(v, places) == round(number, places) for v in recalculated[name])

    # P1 by hand: A-C-F-G = 4 + 8 + 5 + 10 = 27 weeks; normal direct cost 6000+9000+12000+8000+10000+7000+14000 =
    # 66000; best plan 25 weeks: direct 69000 + penalty (25 − 23)·1800 = 3600 → 72600; crashing goes 27 → 20.
    p1 = built["taller2_p1_penalty"]
    assert p1[3].cpm.critical_paths == [["A", "C", "F", "G"]] and p1[3].cpm.project_duration == 27
    assert all(has("taller2_p1_penalty", n) for n in (27, 66000, 25, 69000, 3600, 72600, 20))
    assert [s.project_duration for s in p1[3].crashing.states] == list(range(27, 19, -1))
    assert [s.tab for s in p1[4].steps if s.key.startswith("crash:")][-1].endswith("21-20")
    # P2 by hand: critical path A-C-E-F, t_e = 3 + (3+20+13)/6 + 3 + (1+8+9)/6 = 3 + 6 + 3 + 3 = 15;
    # σ² = (2/6)² + (10/6)² + (2/6)² + (8/6)² = (4 + 100 + 4 + 64)/36 = 43/9; σ = 2.1858; Z = 1/σ → Φ = 0.6763
    p2 = built["taller2_p2_pert"]
    assert p2[3].cpm.project_duration == 15 and p2[3].variance.total_variance == pytest.approx(43 / 9)
    sigma = math.sqrt(43 / 9)
    assert round(sigma, 4) == 2.1858
    assert all(has("taller2_p2_pert", n) for n in (15, 43 / 9, sigma, 0.6763))


def test_progressive_failure_falls_back_to_classic_with_a_note(output_dir, monkeypatch):
    import pm_mcp.service as svc

    def broken(solution, path):
        raise RuntimeError("simulated layout failure")

    monkeypatch.setattr(svc, "write_progressive_workbook", broken)
    result = generate_workbook(taller2(1), "fallback.xlsx")
    assert result.ok and result.layout == "classic" and result.sheets[0] == "Enunciado"
    assert any("simulated layout failure" in n and "classic layout delivered" in n for n in result.notes)
