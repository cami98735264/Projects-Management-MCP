"""AON drawing: crossing minimisation, straight critical path, routed long arcs, deterministic output."""

import json
from pathlib import Path

import pytest

from pm_mcp.domain.models import ProjectDraft
from pm_mcp.ingestion.normalize import normalize_project_input
from pm_mcp.ingestion.requirements_inference import suggest_requirements_from_text
from pm_mcp.output import network_diagram as nd
from pm_mcp.pipeline import solve_project

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


def _cpm(case: str):
    draft = ProjectDraft.model_validate(json.loads((EXAMPLES / case / "project_draft.json").read_text(encoding="utf-8")))
    return solve_project(normalize_project_input(draft).project).cpm


@pytest.mark.parametrize("case, expected", [
    ("taller1", 1),                    # B,C → F,G is a complete bipartite K2,2: one crossing is unavoidable
    ("taller2_problema1_cpm", 0),
    ("taller2_problema2_pert", 0),
    ("pert_ejemplo_4_2", 0),
    ("cpm_compresion_problema_2", 0),
])
def test_minimum_crossings(case, expected):
    assert nd.count_crossings(_cpm(case)) == expected


@pytest.mark.parametrize("case", ["taller1", "taller2_problema1_cpm", "pert_ejemplo_4_2", "cpm_compresion_problema_2"])
def test_critical_path_is_one_straight_line(case):
    cpm = _cpm(case)
    lay = nd.compute_layout(cpm)
    heights = {lay.y[a] for a in cpm.critical_paths[0]}
    assert len(heights) == 1


def test_long_arcs_are_routed_through_free_slots():
    cpm = _cpm("pert_ejemplo_4_2")  # C (column 2) → G (column 4) and H (column 2) → I (column 5)
    lay = nd.compute_layout(cpm)
    for (u, v), chain in lay.chains.items():
        assert len(chain) == lay.layer[v] - lay.layer[u] + 1
        for node in chain[1:-1]:
            assert lay.is_virtual(node)
            real = [n for n in lay.order[lay.layer[node]] if not lay.is_virtual(n)]
            # the routed arc never passes through a node box (box half height 0.8 of a 2.2 pitch → 0.37 pitches)
            assert all(abs(lay.y[node] - lay.y[n]) > 0.37 for n in real)


def test_grid_positions_are_unique_and_rendering_is_deterministic():
    cpm = _cpm("taller1")
    grid = nd.layout(cpm)
    assert len(set(grid.values())) == len(grid) == len(cpm.activities)
    assert nd.render_network_png(cpm) == nd.render_network_png(cpm)


def test_requirements_recognise_normal_duration_and_planning_questions():
    s = suggest_requirements_from_text("a) ¿Cuál es la duración normal del proyecto y su costo?\n"
                                       "b) ¿Qué duración del proyecto debe planear la constructora?")
    outputs = {q.label: q.outputs for q in s.questions}
    assert "EXPECTED_DURATION" in outputs["a"] and "CRASHING" in outputs["b"]


def test_crashing_answer_reads_as_a_decision():
    draft = ProjectDraft.model_validate(json.loads(
        (EXAMPLES / "taller2_problema1_cpm" / "project_draft.json").read_text(encoding="utf-8")))
    answers = {a.label: a for a in solve_project(normalize_project_input(draft).project).answers}
    a, b = answers["a"].answer.split("\n"), answers["b"].answer.split("\n")
    assert a[:3] == ["• Duración normal del proyecto: 30 semanas.",
                     "• Costo directo normal: $69.000 (suma de los costos normales).", "• Ruta crítica: B → E → G"]
    assert b[0] == "• Duración normal: 30 semanas, costo directo $69.000."
    assert b[1].startswith("• Duración recomendada: 25 semanas → costo total mínimo $73.000")
    assert b[2] == "• Para lograrlo se reduce: G 4 semanas, E 1 semana."
    titles = [s.title for s in answers["b"].procedure_steps]
    assert titles[0].startswith("Pendiente de costo") and titles[-1].startswith("Costo total")
    assert "   Paso 4: T = 26, costo directo $72.000 → se reduce E en 1 (+$1.000)" in answers["b"].procedure


# ---------------------------------------------------------------------------------------------- cell drawing

from pm_mcp.output.network_cells import NODE_COLS, NODE_ROWS, plan_cell_network  # noqa: E402
from pm_mcp.service import generate_workbook  # noqa: E402


@pytest.mark.parametrize("case", ["taller1", "taller2_problema1_cpm", "taller2_problema2_pert", "pert_ejemplo_4_2",
                                  "cpm_compresion_problema_2"])
def test_cell_network_geometry(case):
    cpm = _cpm(case)
    net = plan_cell_network(cpm)
    blocks = {(p.row + r, p.col + c) for p in net.nodes.values() for r in range(NODE_ROWS) for c in range(NODE_COLS)}
    assert len(blocks) == len(net.nodes) * NODE_ROWS * NODE_COLS, "node blocks never overlap"
    assert not blocks & set(net.lines), "no line is drawn over a node"
    arrows = [(r, c) for (r, c), lc in net.lines.items() if lc.arrow]
    assert len(arrows) == sum(len(r.successors) for r in cpm.activities), "one arrow head per arc"
    starts = {(p.row + r, p.col) for p in net.nodes.values() for r in range(NODE_ROWS)}
    assert all((r, c + 1) in starts for r, c in arrows), "every arrow head touches its successor's block"
    for (r, c), lc in net.lines.items():  # every line end is matched by its neighbour (no dangling stubs)
        for d, (dr, dc) in {"E": (0, 1), "W": (0, -1), "N": (-1, 0), "S": (1, 0)}.items():
            if d in lc.dirs and not lc.arrow:
                other = net.lines.get((r + dr, c + dc))
                touches_node = (r + dr, c + dc) in blocks
                assert touches_node or (other and {"E": "W", "W": "E", "N": "S", "S": "N"}[d] in other.dirs)
    assert net.crossings == nd.count_crossings(cpm)


def test_critical_arcs_are_double_red_lines():
    net = plan_cell_network(_cpm("taller2_problema1_cpm"))   # B → E → G critical
    crit = [lc for lc in net.lines.values() if lc.critical]
    assert crit and all("═" in lc.text()[0] or "║" in lc.text()[0] or "╗" in lc.text()[0] or "╚" in lc.text()[0]
                        or "╔" in lc.text()[0] or "╝" in lc.text()[0] for lc in crit)
    assert all("═" not in lc.text()[0] for lc in net.lines.values() if not lc.critical)


def test_workbook_network_is_cells_not_pictures(output_dir):
    import zipfile

    import openpyxl
    draft = ProjectDraft.model_validate(json.loads(
        (EXAMPLES / "taller2_problema2_pert" / "project_draft.json").read_text(encoding="utf-8")))
    result = generate_workbook(draft, "red_celdas.xlsx", layout="classic")
    assert result.ok and result.workbook_validation.passed
    with zipfile.ZipFile(result.path) as book:
        assert not [n for n in book.namelist() if n.startswith("xl/media/")]
    sheet = openpyxl.load_workbook(result.path)["Red AON"]
    assert not sheet._images
    formulas = [c.value for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("='CPM'!")]
    assert len(formulas) == 5 * 6, "every node shows t, IC, TC, IL, TL, holgura … linked to the CPM sheet"
    arrows = [c.value for row in sheet.iter_rows() for c in row if isinstance(c.value, str) and c.value.endswith("►")]
    assert len(arrows) == 5 + 2  # 5 arcs + the two legend samples
