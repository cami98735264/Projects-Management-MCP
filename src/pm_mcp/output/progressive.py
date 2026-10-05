"""Progressive, cumulative solution workbook: one sheet per solution step.

Every sheet is the previous sheet plus the new step (highlighted in yellow); the last sheet holds the whole
solution. Nothing is recomputed here: the classic writer (:class:`pm_mcp.output.workbook._Writer`) draws every
sheet through a :class:`~pm_mcp.output.recording.RecordingSheet`, and this module

1. assigns every recorded cell / merge / conditional format to the step in which it appears (the *mask*, from
   the writer's tags and fixed coordinates — see the ``_mask_*`` functions);
2. lays the classic sheets out ONCE as blocks of the final sheet (the "lanes" below), so a block never moves
   between sheets;
3. replays, for step k, every record whose step ≤ k into a fresh xlsxwriter sheet, translating coordinates and
   rewriting the references of formulas and conditional formats (openpyxl Tokenizer, ``$`` kept) so each sheet
   only references itself and only cells that are already visible.

xlsxwriter (not openpyxl) writes the file because the round-trip validator compares cached formula results.

Lanes (column widths are global to a sheet, so blocks that need different widths get different columns):

* rows 1–2: banner of the step; then every block one below the other (:data:`LANE_ORDER`): the problem blocks
  (context, data, questions), the main AON network, the tables (PERT, CPM, probability, crashing, cost curve),
  the Gantt chart, the network of every crashing step and the answers (Resultados);
* the two drawings keep their classic columns from column A (main and step networks share one column structure,
  so the box-drawing arrows still join); every other block spreads each classic column over as many merged
  columns as it needs, so nothing is pushed off screen to the right.
* new content is highlighted with a yellow fill; no emoji anywhere (:func:`_clean`).
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import xlsxwriter
from openpyxl.formula.tokenizer import Token, Tokenizer
from xlsxwriter.utility import xl_col_to_name

from pm_mcp.domain.numbers import number_text
from pm_mcp.i18n import t
from pm_mcp.output.recording import Rec, RecordingSheet
from pm_mcp.output.workbook import GREY, SheetKind as K, WorkbookBuild, _Writer
from pm_mcp.solution import ProjectSolution

NEW_FILL = "#FFF2CC"     # cells that are new in the current sheet
COL_W = 13.7             # uniform column width of the table lane
DEFAULT_W = 8.43         # Excel's default column width
MAX_TAB = 27             # ≤ 31 even after merge.py adds "P1 " / "P12 "
DRAWN = ("net", "step_net")                                       # lanes kept in their own classic columns
LANE_ORDER = ("top", "net", "left", "gantt", "step_net", "bottom")  # reading order, top to bottom
_EMOJI = re.compile("[\U0001F000-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\uFE0F\u200D]")
_BAD_TAB_CHARS = re.compile(r"[\[\]:*?/\\]")
_REF = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")


class ProgressiveError(RuntimeError):
    """The progressive layout could not be built consistently (the service then falls back to classic)."""


@dataclass
class Step:
    key: str
    label: str          # "1.2" or "4"
    title: str
    note: str
    tab: str = ""

    def heading(self, lang: str) -> str:
        return self.label if "." in self.label else f"{t(lang, 'prog.step')} {self.label}"


@dataclass
class _Source:
    name: str                       # classic sheet name (formulas reference it) or a private name
    recs: list[Rec]
    mask: Callable[[Rec], str]


@dataclass
class _Block:
    key: str
    source: _Source
    lane: str                       # top | left | net | step_net | gantt | bottom
    title: str
    max_row: int = 0
    max_col: int = 0
    widths: dict[int, float] = field(default_factory=dict)
    top: int = 0                    # header row; classic row r lands on row top + r
    cols: dict[int, tuple[int, int]] = field(default_factory=dict)  # classic column → (first, last) column
    first_col: int = 1
    last_col: int = 1


@dataclass
class ProgressiveBuild(WorkbookBuild):
    steps: list[Step] = field(default_factory=list)
    first_step: dict[tuple[int, int], int] = field(default_factory=dict)  # final layout cell → step index
    structural: set[tuple[int, int]] = field(default_factory=set)         # banner, block headers, fences
    tab_cells: dict[str, dict[tuple[int, int], int]] = field(default_factory=dict)  # tab → cell → step index
    classic: WorkbookBuild | None = None
    blocks: list[_Block] = field(default_factory=list)                     # final layout of every block


# ---------------------------------------------------------------------------------------------- helpers


def _tab_name(text: str, taken: set[str]) -> str:
    name = _BAD_TAB_CHARS.sub("-", text).strip()[:MAX_TAB].strip()
    base, i = name, 2
    while name in taken:
        suffix = f" ({i})"
        name = base[:MAX_TAB - len(suffix)] + suffix
        i += 1
    taken.add(name)
    return name


def _split_sheet(ref: str) -> tuple[str | None, str]:
    if "!" not in ref:
        return None, ref
    sheet, rest = ref.rsplit("!", 1)
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, rest


def _parse_ref(text: str) -> tuple[bool, int, bool, int]:
    m = _REF.match(text)
    if not m:
        raise ProgressiveError(f"unsupported reference {text!r}")
    col = 0
    for ch in m.group(2).upper():
        col = col * 26 + ord(ch) - 64
    return bool(m.group(1)), col, bool(m.group(3)), int(m.group(4))


def _ref_text(abs_col: bool, col: int, abs_row: bool, row: int) -> str:
    return f"{'$' if abs_col else ''}{xl_col_to_name(col - 1)}{'$' if abs_row else ''}{row}"


def _clean(text: str) -> str:
    """Text without emoji (the workbook must look hand-made): drops pictographs and joins the spaces they leave."""
    return re.sub(r"  +", " ", _EMOJI.sub("", text)).strip() if _EMOJI.search(text) else text


def _text_height(text: str, width: float) -> float:
    return _Writer._text_height(text, width)


# ---------------------------------------------------------------------------------------------- the builder


class _Progressive:
    def __init__(self, solution: ProjectSolution):
        self.s = solution
        self.w = _Writer(solution, io.BytesIO())
        self.classic = self.w.write()
        self.lang = self.w.lang
        self.p = self.w.p
        crash = self.s.crashing if self.w.has(K.CRASHING) else None
        self.last_state = crash.states[-1].step if crash and crash.states else 0
        self.steps = self._plan_steps()
        self.index = {s.key: i for i, s in enumerate(self.steps)}
        self.sources: dict[str, _Source] = {}
        self.blocks = self._plan_blocks()
        self.step_of: dict[tuple[str, int], int] = {}
        for b in self.blocks:
            for rec in b.source.recs:
                key = b.source.mask(rec)
                if key not in self.index:
                    raise ProgressiveError(f"{b.source.name}: no step {key!r} for record {rec}")
                self.step_of[(b.source.name, rec.seq)] = self.index[key]
        self.hist: dict[str, dict[tuple[int, int], list[Rec]]] = {}
        for b in self.blocks:
            h: dict[tuple[int, int], list[Rec]] = {}
            for rec in b.source.recs:
                if rec.op == "cell":
                    h.setdefault((rec.r, rec.c), []).append(rec)
            self.hist[b.source.name] = h
        self.by_name = {b.source.name: b for b in self.blocks}
        self._layout()
        self.retarget = self._sheet_mentions()
        self.extent = {b.source.name: self._extents(b) for b in self.blocks}

    # ------------------------------------------------------------------ mentions of classic sheets in texts

    def _sheet_mentions(self) -> re.Pattern | None:
        """Fixed texts of the classic writer name its sheets ("ver hoja 'Compresión'", "cada nodo enlaza con la hoja
        CPM"). In this layout each classic sheet is a block, so the mention is pointed at the tab where that block
        first appears (every later tab shows it too)."""
        tab_of: dict[str, str] = {}
        for b in self.blocks:
            if b.source.name.startswith("__"):
                continue
            first = min((self.step_of[(b.source.name, rec.seq)] for rec in b.source.recs if rec.op == "cell"),
                        default=None)
            if first is not None:
                tab_of[b.source.name] = self.steps[first].tab
        problem = self.w.names.get(K.PROBLEM)
        if problem and problem not in tab_of:
            tab_of[problem] = self.steps[self.index.get("questions", 0)].tab
        self.tab_of = tab_of
        names = "|".join(re.escape(n) for n in sorted(tab_of, key=len, reverse=True))
        if not names:
            return None
        # 'Name' | hoja Name | sheet Name | the Name sheet — one pass, so a replaced tab is never re-matched
        return re.compile(rf"'(?P<q>{names})'|\b(?P<w>hoja|sheet) (?P<b>{names})\b(?!')|\bthe (?P<t>{names}) sheet\b")

    def _text(self, text: str) -> str:
        text = _clean(text)
        if self.retarget is None:
            return text

        def repl(m: re.Match) -> str:
            if m.group("q"):
                return f"'{self.tab_of[m.group('q')]}'"
            if m.group("b"):
                return f"{m.group('w')} '{self.tab_of[m.group('b')]}'"
            return f"sheet '{self.tab_of[m.group('t')]}'"

        return self.retarget.sub(repl, text)

    # ------------------------------------------------------------------ steps

    def _plan_steps(self) -> list[Step]:
        lang, w = self.lang, self.w
        keys = ["ctx", "data"] + (["questions"] if self.p.questions else [])
        if w.has(K.PERT):
            keys.append("pert")
        if w.has(K.CPM):
            keys += ["net", "fwd", "bwd", "slack"]
        if w.has(K.GANTT):
            keys.append("gantt")
        if w.has(K.PERT) and "pert_variance_first" in w.marks:
            keys.append("variance")
        if w.has(K.PROBABILITY):
            keys.append("prob")
            if "percentile_first" in w.marks:
                keys.append("pct")
        if w.has(K.CRASHING):
            c = self.s.crashing
            keys.append("slopes")
            keys += [f"crash:{st.step}" for st in c.states if st.step > 0]
            if "crash_summary_first" in w.marks:
                keys.append("total")
            if "crash_optimal_first" in w.marks:
                keys.append("decision")
        if w.has(K.COST_CURVE):
            keys.append("curve")
        keys.append("results")

        steps: list[Step] = []
        number = 1
        sub = 0
        taken: set[str] = set()
        states = {st.step: st for st in self.s.crashing.states} if w.has(K.CRASHING) else {}
        for key in keys:
            if key in ("ctx", "data", "questions"):
                sub += 1
                label = f"1.{sub}"
            else:
                number += 1
                label = str(number)
            if key.startswith("crash:"):
                i = int(key.split(":")[1])
                title = t(lang, "prog.title.crash", a=number_text(states[i - 1].project_duration),
                          b=number_text(states[i].project_duration))
                note = t(lang, "prog.note.crash")
                if i == self.last_state:
                    note += " " + t(lang, "prog.note.crash_last")
            else:
                title = t(lang, f"prog.title.{key}")
                note = t(lang, f"prog.note.{key}")
            step = Step(key=key, label=label, title=title, note=note)
            step.tab = _tab_name(f"{label} {title}", taken)
            steps.append(step)
        return steps

    # ------------------------------------------------------------------ masks: classic record → step key

    def _state_step(self, i: int) -> str:
        return "slopes" if i <= 0 else f"crash:{i}"

    def _mask_pert(self, rec: Rec) -> str:
        if rec.has("variance"):
            return "variance"
        if rec.op == "cf" or (rec.op == "cell" and rec.c == 11):
            return "slack"  # the "critical?" column links to the CPM sheet's last column
        return "pert"

    def _mask_cpm(self, rec: Rec) -> str:
        if rec.op == "cf" or rec.has("note"):
            return "slack"
        if rec.op != "cell":
            return "net"
        if rec.r == 2:
            return "fwd"   # project duration = max TC
        if rec.r == 3:
            return "slack"  # critical path
        if rec.r >= 5:
            if rec.c <= 4:
                return "net"
            if rec.c <= 6:
                return "fwd"
            if rec.c <= 8:
                return "bwd"
            return "slack"
        return "net"

    def _mask_network(self, rec: Rec) -> str:
        for tag, step in (("cost_note", "curve"), ("line_crit", "slack"), ("legend_crit", "slack"),
                          ("arc_crit", "slack"), ("slack", "slack"), ("es", "fwd"), ("ef", "fwd"),
                          ("ls", "bwd"), ("lf", "bwd")):
            if rec.has(tag):
                return step if step in self.index else "slack"
        if rec.op == "cf":
            return "slack"
        return "net"

    def _mask_probability(self, rec: Rec) -> str:
        return "pct" if rec.has("pct") else "prob"

    def _mask_crashing(self, rec: Rec) -> str:
        if rec.has("summary"):
            return "total"
        if rec.has("optimal"):
            return "decision"
        if rec.has("stop"):
            return self._state_step(self.last_state)
        state = rec.tag_value("state")
        if state is not None:
            i = int(state)
            if rec.op == "cell" and 5 <= rec.c <= 7:  # what is reduced next: shown with the next state
                return self._state_step(min(i + 1, self.last_state))
            return self._state_step(i)
        return "slopes"

    def _mask_step_network(self, rec: Rec) -> str:
        if rec.has("costpanel"):
            return "total"
        nxt = rec.tag_value("next")
        if nxt is not None:
            return self._state_step(min(int(nxt) + 1, self.last_state))
        state = rec.tag_value("state")
        if state is not None:
            return self._state_step(int(state))
        return "slopes"

    # ------------------------------------------------------------------ generated blocks (1.1 and 1.3)

    def _generated(self, name: str, fill: Callable[[RecordingSheet], None], mask: str) -> _Source:
        g = RecordingSheet(None, lambda spec: spec or {})
        fill(g)
        return _Source(name=name, recs=g.recs, mask=lambda rec, m=mask: m)

    def _fill_context(self, g: RecordingSheet) -> None:
        lang, meta = self.lang, self.p.metadata
        wrap = {"text_wrap": True, "valign": "top"}
        g.set_column(0, 0, 18.7)
        g.set_column(1, 1, 100)
        g.write_string(0, 0, meta.title, {"font_size": 13})
        r = 2
        paragraphs = [line.strip() for line in (meta.context or "").split("\n") if line.strip()]
        for i, para in enumerate(paragraphs):
            if i == 0:
                g.write_string(r, 0, t(lang, "prog.context"), {"valign": "top"})
            g.write_string(r, 1, para, wrap)
            g.set_row(r, min(409.0, _text_height(para, 100)))
            r += 1
        g.write_string(r, 0, t(lang, "time_unit"))
        g.write_string(r, 1, self.w.plural)

    def _fill_questions(self, g: RecordingSheet) -> None:
        head = {"bg_color": GREY, "border": 1, "text_wrap": True, "valign": "top"}
        cell = {"text_wrap": True, "valign": "top", "border": 1}
        g.set_column(0, 0, 10.7)
        g.set_column(1, 1, 100)
        g.write_string(0, 0, t(self.lang, "prog.part"), head)
        g.write_string(0, 1, t(self.lang, "prog.question"), head)
        for i, q in enumerate(self.p.questions):
            g.write_string(1 + i, 0, q.label, {"align": "center", "valign": "top", "border": 1})
            g.write_string(1 + i, 1, q.text, cell)
            g.set_row(1 + i, min(409.0, _text_height(q.text, 100)))

    def _fill_trace(self, g: RecordingSheet) -> None:
        lang = self.lang
        head = {"bg_color": GREY, "border": 1, "text_wrap": True, "valign": "top"}
        for col, width in ((0, 10.7), (1, 10.7), (2, 10.7), (3, 13.7), (4, 13.7), (5, 60.7)):
            g.set_column(col, col, width)
        labels = [t(lang, "activity"), t(lang, "network.tail"), t(lang, "network.head"), t(lang, "predecessors"),
                  t(lang, "network.via"), t(lang, "problem.justification")]
        for i, label in enumerate(labels):
            g.write_string(0, i, label, head)
        for r, tr in enumerate(self.p.arrow_network_trace, start=1):
            values = [tr.activity_id, tr.tail_node, tr.head_node, ", ".join(tr.predecessors) or t(lang, "none"),
                      ", ".join(tr.via_dummies), tr.justification]
            for col, value in enumerate(values):
                if value:
                    g.write_string(r, col, str(value), None)

    # ------------------------------------------------------------------ blocks and layout

    def _classic(self, kind: K, mask: Callable[[Rec], str]) -> _Source:
        return _Source(name=self.w.names[kind], recs=self.w.ws[kind].recs, mask=mask)

    def _plan_blocks(self) -> list[_Block]:
        w, lang = self.w, self.lang
        fixed = lambda key: (lambda rec: key)  # noqa: E731
        blocks = [_Block("ctx", self._generated("__ctx", self._fill_context, "ctx"), "top", t(lang, "prog.block.ctx")),
                  _Block("data", self._classic(K.DATA, fixed("data")), "top", t(lang, "prog.block.data"))]
        if self.p.arrow_network_trace:
            blocks.append(_Block("trace", self._generated("__trace", self._fill_trace, "data"), "top",
                                 t(lang, "prog.block.trace")))
        if self.p.questions:
            blocks.append(_Block("questions", self._generated("__questions", self._fill_questions, "questions"), "top",
                                 t(lang, "prog.block.questions")))
        left = [(K.PERT, "pert", self._mask_pert), (K.CPM, "cpm", self._mask_cpm),
                (K.CPM_DETERMINISTIC, "cpm_det", fixed("slack")), (K.PROBABILITY, "prob", self._mask_probability),
                (K.CRASHING, "crash", self._mask_crashing), (K.COST_CURVE, "curve", fixed("curve"))]
        for kind, key, mask in left:
            if w.has(kind):
                blocks.append(_Block(key, self._classic(kind, mask), "left", t(lang, f"prog.block.{key}")))
        for kind, key, mask in ((K.NETWORK, "net", self._mask_network),
                                (K.STEP_NETWORK, "step_net", self._mask_step_network),
                                (K.GANTT, "gantt", fixed("gantt"))):
            if w.has(kind):
                blocks.append(_Block(key, self._classic(kind, mask), key, t(lang, f"prog.block.{key}")))
        blocks.append(_Block("results", self._classic(K.RESULTS, fixed("results")), "bottom",
                             t(lang, "prog.block.results")))
        for b in blocks:
            for rec in b.source.recs:
                if rec.op == "col":
                    for c in range(rec.c, rec.c2 + 1):
                        b.widths[c] = rec.size if rec.size is not None else DEFAULT_W
                elif rec.op in ("cell", "merge", "cf"):
                    b.max_row = max(b.max_row, rec.r2 or rec.r, rec.r)
                    b.max_col = max(b.max_col, rec.c2 or rec.c, rec.c)
        return blocks

    @staticmethod
    def _span_map(b: _Block, start: int, dest_width: Callable[[int], float]) -> dict[int, tuple[int, int]]:
        """Classic column c → consecutive destination columns whose widths add up to (90 % of) its width."""
        out = {}
        cur = start
        for c in range(1, max(b.max_col, 1) + 1):
            need = 0.9 * b.widths.get(c, DEFAULT_W)
            first, acc = cur, dest_width(cur)
            while acc < need:
                cur += 1
                acc += dest_width(cur)
            out[c] = (first, cur)
            cur += 1
        return out

    def _layout(self) -> None:
        # The drawings (main network and the network of every crashing step) keep their classic columns from
        # column A: they share one column structure (_Writer._network_plans), so they stack without breaking an
        # arrow. Every other block is spread over those widths, and all blocks go one below the other, so the
        # networks are on screen right under the statement instead of thousands of pixels to the right.
        self.widths: dict[int, float] = {}
        for b in self.blocks:
            if b.lane in DRAWN:
                b.cols = {c: (c, c) for c in range(1, b.max_col + 1)}
                for c in range(1, b.max_col + 1):
                    self.widths[c] = max(self.widths.get(c, 0.0), b.widths.get(c, DEFAULT_W))
        width = lambda col: self.widths.get(col, COL_W)  # noqa: E731
        for b in self.blocks:
            if b.lane not in DRAWN:
                b.cols = self._span_map(b, 1, width)
                for _, last in b.cols.values():
                    for col in range(1, last + 1):
                        self.widths.setdefault(col, COL_W)
        self.fences: dict[str, int] = {}  # stacked lanes: no text can spill into another lane
        self.left_width = max(max(last for _, last in b.cols.values())
                              for b in self.blocks if b.lane in ("top", "left"))
        for b in self.blocks:
            b.first_col = min(first for first, _ in b.cols.values())
            b.last_col = max(last for _, last in b.cols.values())
            if b.lane in ("top", "left"):
                b.last_col = max(b.last_col, self.left_width)
        # rows: banner 1–2, then from row 4 every block in reading order, the answers last
        row = 4
        for lane in LANE_ORDER:
            for b in self.blocks:
                if b.lane == lane:
                    b.top = row
                    row = b.top + b.max_row + 2

    def _extents(self, b: _Block) -> dict[tuple[int, int], int]:
        """Last destination column of every classic cell of a spanned block. A plain text cell (no border, fill or
        wrap) is merged as far to the right as its text needs, up to the next cell of the final layout in that
        row (merged cells never let text overflow, so without this a title would be cut at its first column)."""
        out: dict[tuple[int, int], int] = {}
        if b.lane in DRAWN:
            return {key: self._dest(b, *key, end=True)[1] for key in self.hist[b.source.name]}
        occupied: dict[int, list[int]] = {}
        for (r, c) in self.hist[b.source.name]:
            occupied.setdefault(r, []).append(self._dest(b, r, c)[1])
        for rec in b.source.recs:
            if rec.op == "merge":
                for rr in range(rec.r, rec.r2 + 1):
                    occupied.setdefault(rr, []).append(self._dest(b, rec.r, rec.c)[1])
        for (r, c), history in self.hist[b.source.name].items():
            first, last = b.cols[c]
            final = history[-1]
            plain = not ({"border", "text_wrap", "bg_color"} & set(final.spec))
            if final.vtype in ("string", "formula") and plain and isinstance(final.value if final.vtype == "string"
                                                                               else final.cached, str):
                text = self._text(final.value) if final.vtype == "string" else final.cached
                need = len(text) * 1.1 * (final.spec.get("font_size", 11) / 11)
                limit = min([col for col in occupied.get(r, []) if col > first] + [b.last_col + 1]) - 1
                acc = sum(self.widths.get(col, COL_W) for col in range(first, last + 1))
                while acc < need and last < limit:
                    last += 1
                    acc += self.widths.get(last, COL_W)
            out[(r, c)] = last
        return out

    # ------------------------------------------------------------------ references

    def _dest(self, block: _Block, r: int, c: int, end: bool = False) -> tuple[int, int]:
        if c not in block.cols:
            raise ProgressiveError(f"{block.source.name}: column {c} outside the block")
        first, last = block.cols[c]
        return block.top + r, last if end else first

    def _visible(self, block: _Block, r: int, c: int, si: int) -> bool | None:
        """True/False when the classic cell exists (visible yet or not), None when it is blank."""
        history = self.hist[block.source.name].get((r, c))
        if not history:
            return None
        return min(self.step_of[(block.source.name, rec.seq)] for rec in history) <= si

    def _rewrite(self, formula: str, here: _Block, si: int, where: str) -> str:
        tok = Tokenizer(formula)
        for tk in tok.items:
            if tk.type != Token.OPERAND or tk.subtype != Token.RANGE:
                continue
            sheet, rest = _split_sheet(tk.value)
            block = here if sheet is None else self.by_name.get(sheet)
            if block is None:
                raise ProgressiveError(f"{where}: reference to sheet {sheet!r}, which has no block")
            parts = [_parse_ref(p) for p in rest.split(":")]
            if len(parts) == 1:
                ac, c, ar, r = parts[0]
                seen = self._visible(block, r, c, si)
                if seen is None:
                    raise ProgressiveError(f"{where}: {tk.value} points to an empty cell")
                if not seen:
                    raise ProgressiveError(f"{where}: {tk.value} is not visible yet")
                dr, dc = self._dest(block, r, c)
                tk.value = _ref_text(ac, dc, ar, dr)
                continue
            (ac1, c1, ar1, r1), (ac2, c2, ar2, r2) = parts
            for rr in range(min(r1, r2), max(r1, r2) + 1):
                for cc in range(min(c1, c2), max(c1, c2) + 1):
                    if self._visible(block, rr, cc, si) is False:
                        raise ProgressiveError(f"{where}: {tk.value} includes {xl_col_to_name(cc - 1)}{rr}, "
                                               "not visible yet")
            d1 = self._dest(block, r1, c1)
            d2 = self._dest(block, r2, c2, end=c1 != c2)
            tk.value = f"{_ref_text(ac1, d1[1], ar1, d1[0])}:{_ref_text(ac2, d2[1], ar2, d2[0])}"
        return tok.render()

    # ------------------------------------------------------------------ writing

    def _fmt(self, spec: dict[str, Any], new: bool = False):
        full = {**spec, "bold": False}
        if new and "bg_color" not in full:
            full["bg_color"] = NEW_FILL
        key = tuple(sorted((k, str(v)) for k, v in full.items()))
        if key not in self._formats:
            self._formats[key] = self.wb.add_format(full)
        return self._formats[key]

    def _put(self, ws, r: int, c: int, rec: Rec, fmt, text: str | None = None) -> None:
        if rec.vtype == "formula":
            ws.write_formula(r - 1, c - 1, text, fmt, rec.cached)
        elif rec.vtype == "number":
            ws.write_number(r - 1, c - 1, rec.value, fmt)
        elif rec.vtype == "string":
            ws.write_string(r - 1, c - 1, self._text(rec.value), fmt)
        elif fmt is not None:
            ws.write_blank(r - 1, c - 1, None, fmt)

    def write(self, path: Path) -> ProgressiveBuild:
        lang = self.lang
        self.wb = xlsxwriter.Workbook(str(path), {"in_memory": True})
        self.wb.set_properties({"created": datetime(2000, 1, 1), "title": self.p.metadata.title,
                                "author": "pm-scheduling-mcp"})
        self._formats: dict[tuple, Any] = {}
        expected: dict[tuple[str, int, int], Any] = {}
        tab_cells: dict[str, dict[tuple[int, int], int]] = {}
        structural: set[tuple[int, int]] = set()
        first_step: dict[tuple[int, int], int] = {}
        banner_spec = {"font_size": 13, "bg_color": NEW_FILL, "valign": "vcenter"}
        note_spec = {"text_wrap": True, "valign": "top", "bg_color": NEW_FILL}
        for si, step in enumerate(self.steps):
            ws = self.wb.add_worksheet(step.tab)
            cells: dict[tuple[int, int], int] = {}
            for col, width in sorted(self.widths.items()):
                ws.set_column(col - 1, col - 1, width)
            ws.set_row(0, 21)
            ws.set_row(1, 33)
            ws.merge_range(0, 0, 0, self.left_width - 1, _clean(f"{step.heading(lang)} — {step.title}"),
                           self._fmt(banner_spec))
            ws.merge_range(1, 0, 1, self.left_width - 1, _clean(step.note), self._fmt(note_spec))
            structural |= {(1, 1), (2, 1)}
            ws.freeze_panes(2, 0)
            fence_rows: dict[str, set[int]] = {}
            for b in self.blocks:
                name = b.source.name
                shown = [rec for rec in b.source.recs if self.step_of[(name, rec.seq)] <= si]
                if not any(rec.op == "cell" for rec in shown):
                    continue
                latest = max(self.step_of[(name, rec.seq)] for rec in shown)
                last_step = self.steps[latest]
                label = f"{last_step.heading(lang)} — {last_step.title}"
                if b.title != last_step.title:
                    label += f" · {b.title}"
                if latest == si:
                    header, header_fmt = _clean(label), self._fmt({"bg_color": NEW_FILL})
                else:
                    header, header_fmt = _clean(label), self._fmt({})
                if b.last_col > b.first_col:
                    ws.merge_range(b.top - 1, b.first_col - 1, b.top - 1, b.last_col - 1, header, header_fmt)
                else:
                    ws.write_string(b.top - 1, b.first_col - 1, header, header_fmt)
                structural.add((b.top, b.first_col))
                covered: set[tuple[int, int]] = set()
                for rec in shown:
                    if rec.op == "row" and rec.size is not None:
                        ws.set_row(b.top + rec.r - 1, rec.size)
                    elif rec.op == "merge":
                        r1, c1 = self._dest(b, rec.r, rec.c)
                        r2, c2 = self._dest(b, rec.r2, rec.c2, end=True)
                        new = self.step_of[(name, rec.seq)] == si
                        ws.merge_range(r1 - 1, c1 - 1, r2 - 1, c2 - 1, "", self._fmt(rec.spec, new))
                        covered |= {(r, c) for r in range(rec.r, rec.r2 + 1) for c in range(rec.c, rec.c2 + 1)}
                for (r, c), history in self.hist[name].items():
                    chosen = [rec for rec in history if self.step_of[(name, rec.seq)] <= si]
                    if not chosen:
                        continue
                    rec = chosen[-1]
                    rec_step = self.step_of[(name, rec.seq)]
                    fmt = self._fmt(rec.spec, rec_step == si)
                    dr, dc = self._dest(b, r, c)
                    dc_end = self.extent[name][(r, c)]
                    text = None
                    if rec.vtype == "formula":
                        text = self._rewrite(rec.value, b, si, f"{step.tab}!{xl_col_to_name(dc - 1)}{dr}")
                        key = (name, r, c)
                        if key in self.classic.expected_values:
                            expected[(step.tab, dr, dc)] = self.classic.expected_values[key]
                    if dc_end > dc and (r, c) not in covered:
                        ws.merge_range(dr - 1, dc - 1, dr - 1, dc_end - 1, "", fmt)
                    self._put(ws, dr, dc, rec, fmt, text)
                    if rec.vtype != "blank":
                        cells[(dr, dc)] = rec_step
                        first = min(self.step_of[(name, h.seq)] for h in history)
                        first_step[(dr, dc)] = min(first_step.get((dr, dc), first), first)
                    if rec.vtype == "string" and b.lane in self.fences:
                        fence_rows.setdefault(b.lane, set()).add(dr)
                for rec in shown:
                    if rec.op != "cf":
                        continue
                    r1, c1 = self._dest(b, rec.r, rec.c)
                    r2, c2 = self._dest(b, rec.r2, rec.c2, end=True)
                    options = dict(rec.cf)
                    options["format"] = self._fmt(options["format"])
                    if "criteria" in options and str(options["criteria"]).startswith("="):
                        options["criteria"] = self._rewrite(options["criteria"], b, si, f"{step.tab} (format)")
                    rng = f"{xl_col_to_name(c1 - 1)}{r1}:{xl_col_to_name(c2 - 1)}{r2}"
                    ws.conditional_format(rng, options)
            for lane, rows in fence_rows.items():
                col = self.fences[lane]
                for dr in rows:
                    ws.write_string(dr - 1, col - 1, " ")
                    structural.add((dr, col))
            tab_cells[step.tab] = cells
        self.wb.close()
        return ProgressiveBuild(path=str(path), sheet_names=[s.tab for s in self.steps], expected_values=expected,
                                notes=list(self.classic.notes), steps=self.steps, first_step=first_step,
                                structural=structural, tab_cells=tab_cells, classic=self.classic,
                                blocks=self.blocks)


def write_progressive_workbook(solution: ProjectSolution, path: str | Path) -> ProgressiveBuild:
    """Write the progressive workbook to ``path`` (overwriting it) and return what was promised in it."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    return _Progressive(solution).write(path)
