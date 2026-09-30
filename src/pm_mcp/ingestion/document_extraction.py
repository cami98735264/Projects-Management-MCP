"""Thin, deterministic content extraction from the files a user supplies.

This does *not* interpret the exercise. It surfaces everything that is in the file — text in document order,
native tables, spreadsheet cells, and every embedded image (saved to disk with its location and hash) — so the
LLM layer can read text directly and inspect images visually. OCR is optional and always flagged as a
low-confidence hint: the reference files show that tables, formulas and network figures are often *only*
available as images, and OCR misreads them (e.g. letters vs digits).
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import struct
import zipfile
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree as ET

from pydantic import BaseModel, Field

from pm_mcp.domain.errors import IssueCode, IssueSeverity, ValidationIssue

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
R_EMBED = f"{{{NS['r']}}}embed"
R_ID = f"{{{NS['r']}}}id"
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".emf", ".wmf"}


class ExtractedImage(BaseModel):
    index: int
    part: str = Field(description="Path inside the container, e.g. word/media/image2.png")
    location: str = Field(description="Where the image appears, e.g. 'paragraph 12' or \"sheet 'CPM' row 4 col 1\"")
    sha256: str
    size_bytes: int
    width: int | None = None
    height: int | None = None
    saved_path: str | None = None
    ocr_text: str | None = Field(default=None, description="LOW-CONFIDENCE hint only; verify visually")


class ExtractedTable(BaseModel):
    location: str
    rows: list[list[str]]


class ExtractedCell(BaseModel):
    sheet: str
    cell: str
    value: str | int | float | bool
    is_formula: bool = False


class ExtractedDocument(BaseModel):
    path: str
    format: Literal["docx", "xlsx", "pdf", "text", "json", "image"]
    sha256: str
    text_blocks: list[str] = Field(default_factory=list, description="Text in document order; images appear as [IMAGE n]")
    tables: list[ExtractedTable] = Field(default_factory=list)
    sheets: list[str] = Field(default_factory=list)
    cells: list[ExtractedCell] = Field(default_factory=list)
    images: list[ExtractedImage] = Field(default_factory=list)
    json_content: Any | None = None
    warnings: list[ValidationIssue] = Field(default_factory=list)
    guidance: list[str] = Field(default_factory=list)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _image_size(data: bytes) -> tuple[int | None, int | None]:
    if data[:8] == b"\x89PNG\r\n\x1a\n" and len(data) >= 24:
        return struct.unpack(">II", data[16:24])
    try:
        from PIL import Image
        import io

        with Image.open(io.BytesIO(data)) as img:
            return img.size
    except Exception:
        return None, None


def _tesseract_available() -> str | None:
    configured = os.environ.get("PM_MCP_TESSERACT_CMD")
    if configured and Path(configured).exists():
        return configured
    found = shutil.which("tesseract")
    if found:
        return found
    default = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    return str(default) if default.exists() else None


def _ocr(data: bytes, warnings: list[ValidationIssue]) -> str | None:
    command = _tesseract_available()
    try:
        import io

        import pytesseract
        from PIL import Image
    except ImportError:
        warnings.append(ValidationIssue(code=IssueCode.UNSUPPORTED_FORMAT, severity=IssueSeverity.WARNING,
                                        message="OCR requested but pytesseract/Pillow are not installed (extra 'ocr')."))
        return None
    if command is None:
        warnings.append(ValidationIssue(code=IssueCode.UNSUPPORTED_FORMAT, severity=IssueSeverity.WARNING,
                                        message="OCR requested but the tesseract executable was not found (set PM_MCP_TESSERACT_CMD)."))
        return None
    pytesseract.pytesseract.tesseract_cmd = command
    try:
        with Image.open(io.BytesIO(data)) as img:
            return pytesseract.image_to_string(img, config="--psm 6").strip()
    except Exception as exc:
        warnings.append(ValidationIssue(code=IssueCode.UNREADABLE_DOCUMENT, severity=IssueSeverity.WARNING,
                                        message=f"OCR failed: {exc}"))
        return None


class _ImageCollector:
    def __init__(self, output_dir: Path | None, stem: str, ocr: bool, warnings: list[ValidationIssue]):
        self.output_dir = output_dir
        self.stem = stem
        self.ocr = ocr
        self.warnings = warnings
        self.images: list[ExtractedImage] = []

    def add(self, part: str, data: bytes, location: str) -> int:
        index = len(self.images) + 1
        saved = None
        if self.output_dir is not None:
            self.output_dir.mkdir(parents=True, exist_ok=True)
            target = self.output_dir / f"{index:02d}_{Path(part).name}"
            target.write_bytes(data)
            saved = str(target)
        width, height = _image_size(data)
        self.images.append(ExtractedImage(
            index=index, part=part, location=location, sha256=_sha256(data), size_bytes=len(data), width=width,
            height=height, saved_path=saved, ocr_text=_ocr(data, self.warnings) if self.ocr else None))
        return index


def _rels(archive: zipfile.ZipFile, rels_part: str) -> dict[str, str]:
    if rels_part not in archive.namelist():
        return {}
    root = ET.fromstring(archive.read(rels_part))
    return {rel.get("Id"): rel.get("Target") for rel in root.findall("rel:Relationship", NS)}


def _join_part(base_dir: str, target: str) -> str:
    parts: list[str] = [p for p in base_dir.split("/") if p]
    for piece in target.split("/"):
        if piece == "..":
            parts.pop()
        elif piece and piece != ".":
            parts.append(piece)
    return "/".join(parts)


def _extract_docx(archive: zipfile.ZipFile, doc: ExtractedDocument, images: _ImageCollector) -> None:
    rels = _rels(archive, "word/_rels/document.xml.rels")
    body = ET.fromstring(archive.read("word/document.xml")).find("w:body", NS)
    seen_parts: dict[str, int] = {}

    def paragraph_text(p: ET.Element, location: str) -> str:
        out: list[str] = []
        for node in p.iter():
            tag = node.tag.split("}")[1] if "}" in node.tag else node.tag
            if tag == "t" and node.text:
                out.append(node.text)
            elif tag == "tab":
                out.append("\t")
            elif tag in ("br", "cr"):
                out.append("\n")
            elif tag in ("blip", "imagedata"):
                rid = node.get(R_EMBED) or node.get(R_ID)
                target = rels.get(rid)
                if target:
                    part = _join_part("word", target)
                    if part not in seen_parts:
                        seen_parts[part] = images.add(part, archive.read(part), location)
                    out.append(f"[IMAGE {seen_parts[part]}]")
        return "".join(out)

    paragraph_no = 0
    for child in body:
        tag = child.tag.split("}")[1]
        if tag == "p":
            paragraph_no += 1
            text = paragraph_text(child, f"paragraph {paragraph_no}")
            if text.strip():
                doc.text_blocks.append(text)
        elif tag == "tbl":
            table_no = len(doc.tables) + 1
            rows: list[list[str]] = []
            for tr in child.findall("w:tr", NS):
                row: list[str] = []
                for tc in tr.findall("w:tc", NS):
                    texts = [paragraph_text(p, f"table {table_no}") for p in tc.findall(".//w:p", NS)]
                    cell = "\n".join(x for x in texts if x.strip())
                    row.append(cell)
                    doc.text_blocks.extend(x for x in texts if x.strip())
                rows.append(row)
            doc.tables.append(ExtractedTable(location=f"table {table_no}", rows=rows))

    # Images referenced from other parts (e.g. picture bullets in numbering.xml, headers) and orphan media.
    names = set(archive.namelist())
    for name in sorted(names):
        if name.startswith("word/_rels/") and name.endswith(".rels") and name != "word/_rels/document.xml.rels":
            owner = "word/" + Path(name).name[: -len(".rels")]
            for target in _rels(archive, name).values():
                part = _join_part("word", target)
                if Path(part).suffix.lower() in IMAGE_SUFFIXES and part in names and part not in seen_parts:
                    seen_parts[part] = images.add(part, archive.read(part), f"referenced from {owner}")
    for name in sorted(names):
        if name.startswith("word/media/") and name not in seen_parts:
            seen_parts[name] = images.add(name, archive.read(name), "unreferenced media")


def _extract_xlsx(path: Path, archive: zipfile.ZipFile, doc: ExtractedDocument, images: _ImageCollector) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    doc.sheets = wb.sheetnames
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                value = c.value if isinstance(c.value, (str, int, float, bool)) else str(c.value)
                doc.cells.append(ExtractedCell(sheet=ws.title, cell=c.coordinate, value=value,
                                               is_formula=isinstance(value, str) and value.startswith("=")))
                if isinstance(value, str):
                    doc.text_blocks.append(f"[{ws.title}!{c.coordinate}] {value}")

    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    wb_rels = _rels(archive, "xl/_rels/workbook.xml.rels")
    for sheet in workbook.find("s:sheets", NS):
        name = sheet.get("name")
        sheet_part = _join_part("xl", wb_rels[sheet.get(R_ID)])
        sheet_rels = _rels(archive, _join_part(str(Path(sheet_part).parent).replace("\\", "/"),
                                               f"_rels/{Path(sheet_part).name}.rels"))
        for target in sheet_rels.values():
            if "drawings/" not in target:
                continue
            drawing_part = _join_part(str(Path(sheet_part).parent).replace("\\", "/"), target)
            drawing_rels = _rels(archive, _join_part(str(Path(drawing_part).parent).replace("\\", "/"),
                                                     f"_rels/{Path(drawing_part).name}.rels"))
            root = ET.fromstring(archive.read(drawing_part))
            anchors = []
            for anchor in list(root):
                start = anchor.find("xdr:from", NS)
                blip = anchor.find(".//a:blip", NS)
                if start is None or blip is None:
                    continue
                row = int(start.find("xdr:row", NS).text)
                col = int(start.find("xdr:col", NS).text)
                media = _join_part(str(Path(drawing_part).parent).replace("\\", "/"), drawing_rels[blip.get(R_EMBED)])
                anchors.append((row, col, media))
            for row, col, media in sorted(anchors):
                images.add(media, archive.read(media), f"sheet '{name}' row {row + 1} col {col + 1}")


def extract_document_content(path: str | Path, image_output_dir: str | Path | None = None,
                             ocr: bool = False) -> ExtractedDocument:
    path = Path(path)
    warnings: list[ValidationIssue] = []
    if not path.exists():
        return ExtractedDocument(path=str(path), format="text", sha256="", warnings=[ValidationIssue(
            code=IssueCode.UNREADABLE_DOCUMENT, message=f"file not found: {path}")])
    data = path.read_bytes()
    suffix = path.suffix.lower()
    output_dir = Path(image_output_dir) if image_output_dir else None
    images = _ImageCollector(output_dir, path.stem, ocr, warnings)
    fmt = {".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx", ".pdf": "pdf", ".json": "json"}.get(
        suffix, "image" if suffix in IMAGE_SUFFIXES else "text")
    doc = ExtractedDocument(path=str(path), format=fmt, sha256=_sha256(data))

    try:
        if fmt in ("docx", "xlsx"):
            with zipfile.ZipFile(path) as archive:
                if fmt == "docx":
                    _extract_docx(archive, doc, images)
                else:
                    _extract_xlsx(path, archive, doc, images)
        elif fmt == "pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                warnings.append(ValidationIssue(code=IssueCode.UNSUPPORTED_FORMAT,
                                                message="PDF text extraction needs the optional 'pdf' extra (pypdf)."))
            else:
                for number, page in enumerate(PdfReader(path).pages, start=1):
                    doc.text_blocks.append(f"[page {number}] {page.extract_text() or ''}")
        elif fmt == "json":
            doc.json_content = json.loads(data.decode("utf-8"))
        elif fmt == "image":
            images.add(path.name, data, "file")
        elif suffix in (".txt", ".md", ".csv", ""):
            try:
                text = data.decode("utf-8")
            except UnicodeDecodeError:
                text = data.decode("latin-1")
            text = text.replace("\r\n", "\n").replace("\r", "\n")
            doc.text_blocks = [block.strip() for block in text.split("\n\n") if block.strip()]
        else:
            warnings.append(ValidationIssue(code=IssueCode.UNSUPPORTED_FORMAT, message=f"unsupported file type '{suffix}'"))
    except (zipfile.BadZipFile, ET.ParseError, KeyError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        warnings.append(ValidationIssue(code=IssueCode.UNREADABLE_DOCUMENT, message=f"could not read {path.name}: {exc}"))

    doc.images = images.images
    doc.warnings = warnings
    if doc.images:
        doc.guidance.append(
            f"{len(doc.images)} embedded image(s) may contain activity tables, formulas or network figures that are NOT "
            "available as text. Inspect each saved image visually before building the ProjectDraft; OCR text is only a hint.")
        doc.guidance.append(
            "Figures with numbered circles and lettered arrows are activity-on-arrow networks: transcribe each arrow "
            "(tail node, head node, dashed = dummy) and call convert_arrow_network to derive predecessors.")
    if fmt == "xlsx" and not any(c.is_formula for c in doc.cells):
        doc.guidance.append("The workbook has no live formulas; its content is static values and/or images.")
    return doc
